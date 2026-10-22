#!/usr/bin/env python3
"""Build an Android QA DOCX with actual Arabic RTL, English, table & image.

This file is TEST INPUT only, never bundled or exposed as a user template.
Generator is deterministic in content, and contains no personal information.
"""
from io import BytesIO
from pathlib import Path
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
FILE=ROOT/"quality/fixtures/qa_rich_arabic.docx"
TITLE="AQUAWORD_QA_ARABIC_DOCX"
ARABIC="تم فحص اتجاه الفقرة العربية ومحتوى المستند التجريبي"
ENGLISH="Mixed English medical study documentation"
TABLE_TOKEN="AQUAWORD_QA_TABLE_CONTENT"

def rtl(paragraph):
    props=paragraph._p.get_or_add_pPr()
    bidi=props.find(qn("w:bidi"))
    if bidi is None:
        bidi=OxmlElement("w:bidi")
        props.append(bidi)
    bidi.set(qn("w:val"),"1")

def build(path=FILE):
    path=Path(path)
    path.parent.mkdir(parents=True,exist_ok=True)
    doc=Document()
    section=doc.sections[0]
    section.page_width=Cm(21)
    section.page_height=Cm(29.7)
    section.top_margin=Cm(2)
    section.bottom_margin=Cm(2)
    doc.add_heading(TITLE,0)
    line=doc.add_paragraph()
    rtl(line)
    arabic=line.add_run(ARABIC)
    arabic.bold=True
    runprop=arabic._r.get_or_add_rPr()
    rtlprop=OxmlElement("w:rtl")
    rtlprop.set(qn("w:val"),"1")
    runprop.append(rtlprop)
    doc.add_paragraph(ENGLISH)
    mixed=doc.add_paragraph()
    rtl(mixed)
    mixed.add_run("نص عربي ")
    mixed.add_run("NSAIDs / MRI / CT")
    mixed.add_run(" ومصطلحات طبية.")
    table=doc.add_table(rows=2,cols=2)
    table.style="Table Grid"
    table.cell(0,0).text="Medical finding"
    table.cell(0,1).text="وصف"
    table.cell(1,0).text=TABLE_TOKEN
    table.cell(1,1).text="نتيجة طبيعية"
    rtl(table.cell(1,1).paragraphs[0])
    im=Image.new("RGB",(256,120),(237,247,251))
    draw=ImageDraw.Draw(im)
    draw.rectangle((20,24,236,96),outline=(13,108,136),width=5)
    draw.line((25,76,90,40,150,65,230,30),fill=(24,127,159),width=5)
    memory=BytesIO()
    im.save(memory,format="PNG")
    memory.seek(0)
    doc.add_picture(memory,width=Cm(6))
    doc.add_paragraph("AquaWord image retention test")
    footer=section.footer.paragraphs[0]
    footer.text="AquaWord QA — Editable DOCX"
    doc.add_page_break()
    doc.add_heading("Document Page Two",level=1)
    doc.add_paragraph("Second page preserved after Android save")
    doc.save(str(path))
    print(f"Created rich bilingual editable DOCX QA fixture: {path}")
    return path

if __name__=="__main__":
    build()
