#!/usr/bin/env python3
"""Check that genuine Word editing retained the fixture's difficult structures."""
from io import BytesIO
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from docx import Document

from build_rich_fixture import TITLE, ARABIC, ENGLISH, TABLE_TOKEN

W="{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

def validate_document(data: bytes, require_saved_text: str | None = None) -> dict:
    if not data.startswith(b"PK"):
        raise AssertionError("Output is not a DOCX ZIP")
    document=Document(BytesIO(data))
    paragraphs="\n".join(p.text for p in document.paragraphs)
    if TITLE not in paragraphs or ARABIC not in paragraphs or ENGLISH not in paragraphs:
        raise AssertionError("Original bilingual paragraphs are missing after DOCX save")
    if not any(TABLE_TOKEN in c.text for t in document.tables for r in t.rows for c in r.cells):
        raise AssertionError("Original editable table cell was lost")
    if not any("Second page preserved" in p.text for p in document.paragraphs):
        raise AssertionError("Original second document page text was lost")
    if require_saved_text and require_saved_text not in paragraphs:
        raise AssertionError("Android edited text missing from saved DOCX")
    with ZipFile(BytesIO(data)) as z:
        files=z.namelist()
        if not any(name.startswith("word/media/") for name in files):
            raise AssertionError("Embedded image was lost after edit/save")
        root=ET.fromstring(z.read("word/document.xml"))
        if not root.findall(f".//{W}bidi"):
            raise AssertionError("RTL paragraph property missing after DOCX save")
        if not root.findall(f".//{W}tbl"):
            raise AssertionError("Word table XML missing")
    return {"mixed_language":True,"rtl":True,"table":True,"image":True,
            "page_two":True,"saved_edit": bool(require_saved_text)}

if __name__=="__main__":
    from pathlib import Path
    from build_rich_fixture import FILE
    print(validate_document(Path(FILE).read_bytes()))
