#!/usr/bin/env python3
"""Independent real LibreOffice DOCX open/save compatibility gate.

Not a claim about the Android APK. LibreOffice Desktop's headless import/export
is used to detect data loss in four DOCX assets BEFORE packaging them in Android.
Must fail if content disappears, Arabic text is lost or editable tables vanish.
"""
from pathlib import Path
from tempfile import TemporaryDirectory
import subprocess
import json
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
SOURCES=ROOT/"templates/generated"
EXPECTED={"research_en","research_ar","clinical_case","study_notes"}

def signature(file):
    document=Document(str(file))
    paragraphs=[p.text.strip() for p in document.paragraphs if p.text.strip()]
    table_cells=[[cell.text.strip() for row in table.rows for cell in row.cells] for table in document.tables]
    return {"paragraphs":paragraphs, "table_cells":table_cells,
            "headings":[p.text.strip() for p in document.paragraphs
                        if p.style is not None and p.style.name.startswith("Heading")]}


def run_test():
    paths=list(SOURCES.glob("*.docx"))
    if {p.stem for p in paths} != EXPECTED:
        raise AssertionError("Exactly four known AquaWord DOCX assets required")
    with TemporaryDirectory(prefix="aquaword-libreoffice-test-") as directory:
        tmp=Path(directory)
        profile=(tmp/"lo-profile").as_uri()
        command=["soffice","-env:UserInstallation="+profile,"--headless",
                 "--convert-to","docx:Office Open XML Text",
                 "--outdir",str(tmp),*(str(path) for path in paths)]
        subprocess.run(command, check=True, timeout=150, capture_output=True, text=True)
        report=[]
        for original in sorted(paths):
            converted=tmp/original.name
            if not converted.is_file() or converted.stat().st_size<1500:
                raise AssertionError(f"LibreOffice could not round-trip {original.name}")
            a,b=signature(original),signature(converted)
            if len(b["paragraphs"])<len(a["paragraphs"])-1:
                raise AssertionError(f"Text blocks vanished from {original.name}")
            expected_sections=set(a["headings"])
            actual_sections=set(b["headings"])
            if not expected_sections.issubset(actual_sections):
                raise AssertionError(f"Lost heading content in {original.name}: {expected_sections-actual_sections}")
            if len(a["table_cells"]) != len(b["table_cells"]):
                raise AssertionError(f"Editable tables were lost from {original.name}")
            if original.stem=="research_ar":
                original_arabic="".join(a["paragraphs"])
                converted_arabic="".join(b["paragraphs"])
                arabic_chars=sum('\u0600'<=c<='\u06ff' for c in converted_arabic)
                if arabic_chars<50:
                    raise AssertionError("Arabic text was lost on DOCX re-save")
                if "البحث الأكاديمي" not in converted_arabic:
                    raise AssertionError("Arabic document title changed unexpectedly")
            report.append({"name":original.name,
                           "size_input":original.stat().st_size,
                           "size_roundtrip":converted.stat().st_size,
                           "heading_count":len(b["headings"]),
                           "table_count":len(b["table_cells"]),
                           "status":"PASS: Desktop LibreOffice import/export"})
        print(json.dumps({"result":"PASS","scope":"Desktop LibreOffice DOCX roundtrip only",
                          "assets":report},ensure_ascii=False,indent=2))


if __name__=="__main__":
    run_test()
