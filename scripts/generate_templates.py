#!/usr/bin/env python3
"""Create bundled editable DOCX templates with python-docx (MIT build-time dependency).

Output DOCX documents are AquaWord assets. The python-docx library itself is not
included in the APK. Source templates intentionally contain no real patient data.
"""
from __future__ import annotations

from pathlib import Path
import zipfile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DEST = ROOT / "templates" / "generated"
TEMPLATES = {
    "research_en": "Research Paper · English",
    "research_ar": "بحث أكاديمي · عربي",
    "clinical_case": "Clinical Case Report",
    "study_notes": "High-Yield Study Notes",
}


def direction(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ppr = paragraph._p.get_or_add_pPr()
    bidi = ppr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        ppr.append(bidi)
    bidi.set(qn("w:val"), "1")
    for run in paragraph.runs:
        rpr = run._r.get_or_add_rPr()
        rtl = rpr.find(qn("w:rtl"))
        if rtl is None:
            rtl = OxmlElement("w:rtl")
            rpr.append(rtl)
        rtl.set(qn("w:val"), "1")


def blank(label: str, document: Document, rtl: bool = False):
    p = document.add_paragraph(label + " " + ("_____________________" if not rtl else "________________"))
    if rtl:
        direction(p)


def paragraph(doc: Document, text: str, rtl: bool = False, style: str | None = None):
    p = doc.add_paragraph(text, style=style)
    if rtl:
        direction(p)
    return p


def heading(doc: Document, text: str, level: int = 1, rtl: bool = False):
    p = doc.add_heading(text, level=level)
    if rtl:
        direction(p)
    return p


def comparison_table(doc: Document, labels: tuple[str, str, str], rtl: bool = False):
    table = doc.add_table(rows=2, cols=3)
    table.style = "Table Grid"
    for i, label in enumerate(labels):
        cell = table.rows[0].cells[i]
        cell.text = label
        for run in cell.paragraphs[0].runs:
            run.bold = True
        if rtl:
            direction(cell.paragraphs[0])
        table.rows[1].cells[i].text = "—"
        if rtl:
            direction(table.rows[1].cells[i].paragraphs[0])
    return table


def document_base(arabic: bool = False) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.1)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial" if arabic else "Aptos"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(35, 48, 65)
    normal.paragraph_format.space_after = Pt(7)
    for name, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 12)):
        st = doc.styles[name]
        st.font.name = "Arial" if arabic else "Aptos Display"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(16, 98, 137)
    # Editable PAGE field, not a hard-coded number.
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("الصفحة " if arabic else "Page ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)
    if arabic:
        direction(footer)
    doc.core_properties.author = "AquaWord"
    doc.core_properties.keywords = "Editable document template"
    return doc


def intro(doc: Document, text: str, rtl: bool = False):
    p = doc.add_paragraph(text, style="Title")
    if rtl:
        direction(p)


def research_english() -> Document:
    d = document_base()
    intro(d, "Research Paper")
    for label in ("Working title:", "Author:", "Institution:", "Submission date:"):
        blank(label, d)
    for section, prompt in [
        ("Abstract", "Write a structured summary of background, methods, results, and conclusion."),
        ("Keywords", "Add 3–6 relevant keywords."),
        ("1. Introduction", "Introduce the topic, existing knowledge, research gap, and objective."),
        ("2. Research question and hypothesis", "Specify the question, endpoints, and hypothesis."),
        ("3. Methods", "Describe design, setting, sampling, measures, and analysis plan."),
        ("4. Results", "Report evidence without interpreting it; use tables and figures where needed."),
        ("5. Discussion", "Interpret findings, compare literature, and explain implications."),
        ("6. Limitations", "Describe sources of bias and uncertainty."),
        ("7. Conclusion", "Summarize key findings without unsupported claims."),
        ("References", "Insert verified citations. For APA 7, check author/year/title/source/DOI and hanging indent."),
        ("Appendix", "Add supporting materials if necessary."),
    ]:
        heading(d, section)
        paragraph(d, prompt)
        if section == "4. Results":
            comparison_table(d, ("Outcome", "Value", "Interpretation"))
    return d


def research_arabic() -> Document:
    d = document_base(True)
    intro(d, "البحث الأكاديمي", True)
    for label in ("عنوان البحث:", "اسم الباحث:", "الجامعة:", "التاريخ:"):
        blank(label, d, True)
    for title, text in [
        ("الملخص", "اكتب ملخصًا يوضح خلفية البحث والمنهجية والنتائج والاستنتاج."),
        ("الكلمات المفتاحية", "اكتب الكلمات المفتاحية المناسبة لموضوع البحث."),
        ("١. المقدمة", "عرّف المشكلة وأهميتها والفجوة البحثية."),
        ("٢. أسئلة البحث والأهداف", "حدّد الأسئلة والفرضيات والأهداف بوضوح."),
        ("٣. منهجية البحث", "صف تصميم الدراسة، العينة، الأدوات وأساليب التحليل."),
        ("٤. النتائج", "قدّم النتائج، والجداول والأشكال عند الحاجة."),
        ("٥. المناقشة", "فسّر النتائج وقارنها بالدراسات السابقة."),
        ("٦. القيود", "ناقش محدوديات البحث واحتمالات التحيّز."),
        ("٧. الخاتمة", "اكتب الخلاصة والتوصيات المدعومة بالنتائج."),
        ("المراجع", "أدرج المراجع الموثقة فقط وفق نظام الاستشهاد المطلوب."),
        ("الملاحق", "أضف الملاحق والبيانات الداعمة عند الحاجة."),
    ]:
        heading(d, title, rtl=True)
        paragraph(d, text, True)
    return d


def clinical_case() -> Document:
    d = document_base()
    intro(d, "Clinical Case Report")
    paragraph(d, "Use de-identified data only. Remove all patient identifiers and obtain consent according to institutional policy.")
    for title, body in [
        ("Case title", "Insert a descriptive, non-identifying clinical title."),
        ("Abstract", "Briefly summarize the presentation, workup, treatment, and outcome."),
        ("Patient information", "Age range, relevant background, and anonymized history. Do not include names, dates of birth, or IDs."),
        ("Chief complaint", "Describe the primary concern."),
        ("History of present illness", "Describe chronology, relevant symptoms, and negatives."),
        ("Medical and medication history", "Include relevant diagnoses, treatments, allergies, and social context."),
        ("Examination", "Document measured vital signs and focused physical examination findings."),
        ("Investigations", "Laboratory studies, imaging, differential workup and objective results."),
        ("Differential diagnosis", "List plausible alternatives and discriminating findings."),
        ("Management", "Document decisions, procedures, and therapeutic approach."),
        ("Follow-up and outcome", "Report time course, complications, response and follow-up."),
        ("Discussion", "Explain clinical reasoning and compare with peer-reviewed evidence."),
        ("Learning points", "Summarize three practical, evidence-supported takeaways."),
        ("References", "Add verifiable references; use Vancouver numbering if requested."),
        ("Consent and ethics", "Document informed consent/approval requirements without identifying details."),
    ]:
        heading(d, title)
        paragraph(d, body)
        if title == "Investigations":
            comparison_table(d, ("Test / imaging", "Result", "Clinical relevance"))
    return d


def study_notes() -> Document:
    d = document_base()
    intro(d, "High-Yield Study Notes")
    for label in ("Course:", "Chapter:", "Learning objective:", "Exam date:"):
        blank(label, d)
    for title, body in [
        ("Big picture", "Summarize the topic in 3–5 sentences."),
        ("Core vocabulary", "Term — contextual definition — example."),
        ("Mechanism / pathophysiology", "Write the causal explanation before memorizing details."),
        ("High-yield facts", "Record exam-relevant findings and common pitfalls."),
        ("Comparison table", "Use Insert → Table to compare similar conditions or concepts."),
        ("Clinical applications", "Work through a realistic case and justify each decision."),
        ("Recall questions", "Q1: ...\nQ2: ...\nQ3: ..."),
        ("MCQs and explanations", "Write questions with one best answer, rationale, and why distractors are incorrect."),
        ("Mistake notebook", "Mistake → reason → corrected concept → next review date."),
        ("References and source pages", "For every claim, note the textbook/source title and page."),
    ]:
        heading(d, title)
        paragraph(d, body)
        if title == "Comparison table":
            comparison_table(d, ("Feature", "Concept A", "Concept B"))
    return d


BUILDERS = {
    "research_en": research_english,
    "research_ar": research_arabic,
    "clinical_case": clinical_case,
    "study_notes": study_notes,
}


def generate() -> None:
    DEST.mkdir(parents=True, exist_ok=True)
    for key, factory in BUILDERS.items():
        file = DEST / (key + ".docx")
        factory().save(str(file))
        if not zipfile.is_zipfile(file):
            raise RuntimeError(f"Not a DOCX package: {file}")
        with zipfile.ZipFile(file) as z:
            if z.testzip() is not None:
                raise RuntimeError(f"Corrupted Word package: {file}")
        print(f"Created {file.name} ({file.stat().st_size} bytes)")


if __name__ == "__main__":
    generate()
