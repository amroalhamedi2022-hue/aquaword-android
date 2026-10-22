"""Test Android edit/save gate accepts only an actual editable DOCX package."""
import importlib.util
from io import BytesIO
from pathlib import Path
import unittest

from docx import Document

FILE=Path(__file__).resolve().parents[1]/"scripts/android_edit_roundtrip.py"
SPEC=importlib.util.spec_from_file_location("android_roundtrip", FILE)
RT=importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(RT)

def word_document(text):
    from tempfile import TemporaryDirectory
    from build_rich_fixture import build
    with TemporaryDirectory() as directory:
        original=build(Path(directory)/"rich-qa.docx")
        document=Document(str(original))
        document.add_paragraph(text)
        f=BytesIO()
        document.save(f)
        return f.getvalue()


class AndroidEditRoundTripTests(unittest.TestCase):
    def test_actual_docx_with_inserted_marker(self):
        self.assertTrue(RT.check_document(word_document(RT.SENTINEL)))

    def test_image_or_table_loss_is_detected_even_with_saved_text(self):
        payload=word_document(RT.SENTINEL)
        original=Document(BytesIO(payload))
        first_table=original.tables[0]._element
        first_table.getparent().remove(first_table)
        out=BytesIO()
        original.save(out)
        with self.assertRaisesRegex(AssertionError, "table"):
            RT.check_document(out.getvalue())

    def test_original_unchanged_docx_must_fail(self):
        self.assertFalse(RT.check_document(word_document("Original unchanged")))

    def test_plain_log_cannot_be_faked_as_docx(self):
        with self.assertRaises(Exception):
            RT.check_document(b"PASS: marker " + RT.SENTINEL.encode())

    def test_qa_calls_actual_uno_commands(self):
        self.assertIn(".uno:GoToEndOfDoc", RT.EXPRESSION)
        self.assertIn(".uno:InsertText", RT.EXPRESSION)
        self.assertIn(".uno:Save", RT.EXPRESSION)
        self.assertIn(RT.SENTINEL, RT.EXPRESSION)


if __name__=="__main__":
    unittest.main()
