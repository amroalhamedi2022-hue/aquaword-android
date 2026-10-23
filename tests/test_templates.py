"""Acceptance checks for AquaWord's offline editable Word templates."""
from pathlib import Path
import sys
import unittest
from zipfile import ZipFile
from xml.etree import ElementTree
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from generate_templates import BUILDERS  # noqa: E402
from apply_templates import ASSETS  # noqa: E402

FILES = ROOT / "templates" / "generated"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class TemplateContracts(unittest.TestCase):
    def test_same_four_templates_in_generator_and_android(self):
        self.assertEqual(set(BUILDERS), set(ASSETS))
        self.assertEqual(set(x.stem for x in FILES.glob("*.docx")), set(BUILDERS) | {"blank"})

    def test_blank_new_document_is_real_empty_docx(self):
        path = FILES / "blank.docx"
        self.assertTrue(path.is_file())
        with ZipFile(path) as archive:
            self.assertIsNone(archive.testzip())
            self.assertIn("word/document.xml", archive.namelist())
        d = Document(path)
        self.assertEqual([p.text for p in d.paragraphs if p.text.strip()], [])
        self.assertEqual(len(d.tables), 0)
        self.assertEqual(round(d.sections[0].page_width.cm, 1), 21.0)

    def test_office_openxml_package_and_content(self):
        for key in BUILDERS:
            with self.subTest(template=key):
                path = FILES / (key + ".docx")
                self.assertTrue(path.is_file())
                with ZipFile(path) as zipdoc:
                    self.assertIsNone(zipdoc.testzip())
                    self.assertIn("word/document.xml", zipdoc.namelist())
                    self.assertIn("[Content_Types].xml", zipdoc.namelist())
                    self.assertFalse(any(p.lower().endswith("vbaproject.bin")
                                         for p in zipdoc.namelist()))
                doc = Document(path)
                self.assertGreater(len(doc.paragraphs), 10)
                self.assertTrue(any(p.style.name.startswith("Heading")
                                    for p in doc.paragraphs))
                self.assertEqual(round(doc.sections[0].page_width.cm, 1), 21.0)
                self.assertEqual(round(doc.sections[0].page_height.cm, 1), 29.7)

    def test_editable_tables_and_page_numbers(self):
        for template in ("research_en", "clinical_case", "study_notes"):
            with self.subTest(template=template):
                path=FILES/(template+".docx")
                d=Document(path)
                self.assertGreaterEqual(len(d.tables),1)
                self.assertEqual(len(d.tables[0].columns),3)
                self.assertGreaterEqual(len(d.tables[0].rows),2)
        for key in BUILDERS:
            with ZipFile(FILES/(key+".docx")) as archive:
                footerfiles=[p for p in archive.namelist() if p.startswith("word/footer") and p.endswith(".xml")]
                self.assertTrue(footerfiles,key)
                self.assertIn(b'PAGE',archive.read(footerfiles[0]),key)

    def test_arabic_is_true_rtl_not_only_right_aligned(self):
        path = FILES / "research_ar.docx"
        with ZipFile(path) as archive:
            root = ElementTree.fromstring(archive.read("word/document.xml"))
            bidi = root.findall(".//" + W + "bidi")
            rtl = root.findall(".//" + W + "rtl")
            self.assertGreaterEqual(len(bidi), 12)
            self.assertGreaterEqual(len(rtl), 12)
            self.assertTrue(all(x.get(W + "val") == "1" for x in bidi))
        text = "\n".join(p.text for p in Document(path).paragraphs)
        self.assertIn("منهجية البحث", text)
        self.assertIn("المراجع", text)

    def test_clinical_template_data_safety_and_scope(self):
        text = "\n".join(p.text for p in Document(FILES / "clinical_case.docx").paragraphs)
        for heading in ("Differential diagnosis", "Consent and ethics", "Investigations"):
            self.assertIn(heading, text)
        self.assertIn("Use de-identified data only", text)
        self.assertNotIn("Patient name:", text)

    def test_study_notes_has_recall_and_references(self):
        text = "\n".join(p.text for p in Document(FILES / "study_notes.docx").paragraphs)
        self.assertIn("High-yield facts", text)
        self.assertIn("MCQs and explanations", text)
        self.assertIn("Mistake notebook", text)
        self.assertIn("References and source pages", text)


if __name__ == "__main__":
    unittest.main()
